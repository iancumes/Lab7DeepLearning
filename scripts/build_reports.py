"""PDF y Word con contenido idéntico, tablas nativas y resultados observados.

Ejecutar con Python del runtime de documentos de Codex, después del entrenamiento.
"""
from pathlib import Path
import json, csv, html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,Image
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]; ART=ROOT/'artifacts'; OUT=ROOT/'entregables'
URL='https://github.com/iancumes/Lab7DeepLearning'
def read(name): return json.loads((ART/name).read_text(encoding='utf-8'))
def percent(x): return 'N/D' if x is None else f'{100*x:.2f}%'
def number(x): return f'{x:,}'.replace(',',' ')
def rho(x): return 'N/D' if x is None else f'{x:.5f}'
def mins(x): return f'{x/60:.1f}'
def p(text): return ('p',text)
def h(text): return ('h',text)
def table(headers,rows): return ('table',headers,rows)
def build_content():
    assert read('completion.json')['complete']
    corpus=read('corpus.json'); cfgs=read('configurations.json'); best=read('best_sgns.json')
    runs={c['name']:read(f"runs/{c['name']}.json") for c in cfgs}
    intrinsic=read('evaluations/intrinsic.json'); parallel=read('evaluations/parallelism.json')
    personal=read('evaluations/personal_analogies.json'); selected=read('classification_selection.json')
    test=read('evaluations/classification_test.json'); hw=read('hardware.json'); metas=read('corpus_configurations.json')
    full=[r for r in test if r['fraction']==1]; selected_full={r['initialization']:r for r in selected['selected'] if r['fraction']==1}
    winner=max(full,key=lambda r:r['metrics']['f1_macro'])
    total_seconds=sum(r['training_seconds'] for rows in runs.values() for r in rows)
    pages=[]
    pages.append([
      h('Investigación y preprocesamiento'),
      p(f"Se implementó SGNS propio y se comparó con gensim y GloVe en analogías, similitud léxica y clasificación AG News. El mejor clasificador observado con todos los datos fue {winner['name']}, con F1 macro {percent(winner['metrics']['f1_macro'])} en test. El notebook conserva resultados extensos, código comentado y manifiestos reproducibles."),
      p('CBOW predice el centro desde el contexto; Skip-Gram predice contexto desde el centro. Negative sampling discrimina pares reales frente a ruido unigram elevado a 0.75. GloVe ajusta productos de vectores a logaritmos de coocurrencias globales ponderadas. Los dos métodos capturan relaciones distribucionales, pero asignan un único vector por palabra y pueden reproducir sesgos.'),
      p('nn.Embedding representa una tabla V × d. from_pretrained permite inicialización y congelamiento. EmbeddingBag reduce cada documento por mean, sum o max; offsets indica comienzos dentro de IDs concatenados. Aquí se usa mean, padding_idx=0 y una bolsa por noticia.'),
      p('El tokenizador propio usa minúsculas, palabras Unicode, apóstrofos internos, números y puntuación; repara marcadores @ y elimina especiales y signos de encabezados. Se comparó con NLTK Treebank en 12 oraciones: las contracciones y posesivos explican diferencias. No se aplicó stemming ni eliminación de stopwords.'),
      table(['Split','Artículos','Líneas','Tokens normalizados'],[[s,number(corpus['splits'][s]['articles']),number(corpus['splits'][s]['lines']),number(corpus['splits'][s]['normalized_tokens'])] for s in ['train','validation','test']]),
      p(f"Se seleccionaron {number(corpus['selected_articles'])} artículos completos del entrenamiento, con semilla 23236, hasta {number(corpus['selected_tokens'])} tokens. Los corpus 25/50/100% son prefijos anidados que terminan en artículos completos. Zipf: los top 10/1000/30000 cubren {percent(corpus['zipf_coverage']['10'])}/{percent(corpus['zipf_coverage']['1000'])}/{percent(corpus['zipf_coverage']['30000'])} de tokens."),
      table(['min_count','Vocabulario seleccionado','Tokens excluidos'],[[r['min_count'],number(r['vocabulary_size']),percent(r['unknown_token_fraction'])] for r in corpus['thresholds']]),
      p('La preparación preserva fronteras de línea y bloques de hasta 1000 tokens. Gensim y SGNS reciben el mismo archivo de tokens elegibles, comprobado por hash y vocabulario. Los especiales y palabras bajo min_count no generan pares de entrenamiento.')])
    exp=[]
    for c in cfgs:
        chosen=max(runs[c['name']],key=lambda r:(r['wordsim']['spearman'] or -2,r['analogies']['total']['accuracy'] or 0))
        exp.append([c['name'],f"{c['dim']}/{c['window']}/{c['negatives']}",f"{c['fraction']*100:g}",chosen['epoch'],rho(chosen['wordsim']['spearman']),percent(chosen['analogies']['total']['accuracy']),mins(sum(r['training_seconds'] for r in runs[c['name']]))])
    base=runs['base100'][0]; meta=metas[0]
    pages.append([
      h('Entrenamiento y selección'),
      p('El SGNS propio usa dos tablas nn.Embedding dispersas. Minimiza −log σ(v·u) − Σ log σ(−v·n), preservando los vectores de entrada. Genera pares por bloques, sin cruzar contextos; rechaza negativos iguales al contexto positivo. La suma del minibatch guía SGD y la pérdida publicada se promedia por par.'),
      p('Cada configuración tiene tres epochs, minibatch de 4096 pares, min_count=5, subsampling t=10⁻⁴ y SGD con learning rate lineal 0.025→0.0001. La ventana es fija. Se guardan checkpoint, pérdida, analogías semánticas/sintácticas, WordSim, vecinos de seis palabras, tiempos y memoria por epoch. La selección usa WordSim sin redondear sobre pares comunes y accuracy de analogías como desempate; SimLex queda reservado.'),
      table(['Configuración','d/w/k','Corpus %','Epoch','ρ WordSim','Acc analogías','Min total'],exp),
      p(f"Selección: {best['config']['name']}, epoch {best['epoch']}, ρ WordSim={rho(best['wordsim']['spearman'])}. En la base, el subsampling redujo {number(meta['pairs_before_subsampling'])} pares a {number(base['pairs'])} en la primera epoch; conservó {number(base['retained_tokens'])} tokens elegibles. La pérdida media base pasó de {runs['base100'][0]['loss']:.3f} a {runs['base100'][-1]['loss']:.3f}."),
      p('El descarte observado de the/of/and en la primera epoch fue '+', '.join(f"{w}: {percent(s['observed_discard_fraction'])}" for w,s in base['subsampling'].items())+'.'),
      p(f"Hardware medido: {hw['gpu'] or hw['processor']}, PyTorch {hw['torch']}, Python {hw['python']}; tiempo SGNS acumulado de las siete configuraciones {mins(total_seconds)} minutos. La memoria GPU máxima asignada a tensores durante entrenamiento fue {max(r['gpu_peak_bytes'] or 0 for rows in runs.values() for r in rows)/1024**2:.1f} MiB. Estos tiempos excluyen descarga y evaluación."),
      p('Gensim usa el mismo corpus/vocabulario y parámetros equivalentes, cuatro workers y ventana fija. Sus actualizaciones Cython asíncronas difieren de los minibatches; descarta colisiones de negativos en vez de redibujarlas. Se registra cada epoch, sin exigir vectores numéricamente iguales.'),
      p('GloVe es exactamente glove-wiki-gigaword-100: 400000 palabras, 100d y 6 mil millones de tokens. El artículo informa dos Intel Xeon E5-2658 de 2.1 GHz, 85 min de coocurrencias en un hilo y 14 min por iteración de 300d en 32 cores. Son mediciones publicadas de otra dimensión; aquí no se entrenó GloVe.')])
    examples=[]
    for name,rows in personal.items():
        valid=[r for r in rows if not r['oov']]
        hits=[r for r in valid if r['rank']==1]
        failures=sorted([r for r in valid if r['rank']>1],key=lambda r:r['rank'],reverse=True)
        # Un ejemplo por modelo; las 54 consultas completas están en el notebook.
        for label,pool in [('Acierto' if hits else 'Fallo',hits or failures)]:
            if pool:
                r=pool[0]
                examples.append([name,label,f"{r['b']} − {r['a']} + {r['c']}",r['expected'],r['top5'][0][0],r['rank']])
    covered=intrinsic['sgns']['three_cos_add']['total']
    labels=['Capitales comunes','Capitales mundo','Monedas','Ciudades/estados','Familia',
            'Adj.→adverbio','Opuestos','Comparativos','Superlativos','Gerundios',
            'Gentilicios','Pasado','Plurales','Verbos plurales']
    categoryrows=[]
    for index,label in enumerate(labels):
        values=[]
        for model in ['sgns','gensim','glove']:
            values.append('/'.join('N/D' if intrinsic[model][method]['categories'][index]['accuracy'] is None else f"{100*intrinsic[model][method]['categories'][index]['accuracy']:.1f}" for method in ['three_cos_add','three_cos_mul']))
        categoryrows.append([label]+values)
    pages.append([
      h('Aritmética vectorial y geometría'),
      p('analogia(a,b,c,k) implementa 3CosAdd con vectores individuales normalizados y búsqueda por coseno de b−a+c. Se excluyen las consultas a/b/c y se verificó coincidencia de top-5 y similitudes con gensim. 3CosMul usa el cociente de productos de cosenos desplazados a [0,1], con epsilon=10⁻⁶. Las consultas excluidas impiden respuestas triviales.'),
      p(f"Los tres embeddings se evalúan sobre las mismas {number(covered['evaluated'])} preguntas cubiertas de {number(covered['total'])}, en 14 categorías, con exactamente 30000 candidatos compartidos: cobertura {percent(covered['coverage'])}. Accuracy usa solo preguntas cubiertas; la cobertura debe acompañarla. Los resultados por categoría están en el notebook."),
      table(['Modelo','Semántica Add','Sintáctica Add','Total Add','Total Mul','Paralelismo'],[[name,percent(r['three_cos_add']['semantic']['accuracy']),percent(r['three_cos_add']['syntactic']['accuracy']),percent(r['three_cos_add']['total']['accuracy']),percent(r['three_cos_mul']['total']['accuracy']),rho(parallel[name]['mean_cosine'])] for name,r in intrinsic.items()]),
      p('Se probaron 18 analogías personales de seis tipos, con top-5, similitudes, rango correcto, OOV y búsqueda con/sin exclusión. Ejemplos observados:'),
      table(['Modelo','Resultado','Consulta','Esperado','Top 1','Rango'],examples),
      table(['Categoría','SGNS Add/Mul %','Gensim Add/Mul %','GloVe Add/Mul %'],categoryrows),
      p('El coseno entre b−a y d−c mide paralelismo de relaciones equivalentes; un valor alto no garantiza acertar el ranking. El notebook incluye t-SNE de aproximadamente 500 palabras por modelo, con grupos definidos previamente, coseno, perplexity 30, 1500 iteraciones y semilla 23236. Sus ejes y distancias globales no son medidas semánticas.')])
    fractionrows=[]
    for f in [.01,.1,.5,1.0]:
        values={r['initialization']:r for r in test if r['fraction']==f}
        fractionrows.append([f'{100*f:g}%']+[percent(values[m]['metrics']['f1_macro']) for m in ['tfidf','random','sgns','gensim','glove']])
    fullmetrics=[]
    for r in full:
        chosen=selected_full[r['initialization']]
        fullmetrics.append([r['name'],percent(r['metrics']['accuracy']),percent(r['metrics']['precision_macro']),percent(r['metrics']['recall_macro']),percent(r['metrics']['f1_macro']),number(chosen['trainable_parameters']),mins(chosen['training_seconds'])])
    confusion=max(full,key=lambda r:sum(np for i,row in enumerate(r['confusion_matrix']) for j,np in enumerate(row) if i!=j))
    errors=[(count,i,j) for i,row in enumerate(confusion['confusion_matrix']) for j,count in enumerate(row) if i!=j]
    count,i,j=max(errors)
    label=['World','Sports','Business','Sci/Tech']
    pages.append([
      h('Clasificación de noticias'),
      p('AG News se separó en 108000 ejemplos de entrenamiento y 12000 de validación estratificada, sin solapamiento de índices; test oficial contiene 7600. Las fracciones 1/10/50/100% usan 1080/10800/54000/108000 ejemplos anidados por clase. Cada vocabulario y TF-IDF se aprende únicamente desde la fracción correspondiente.'),
      p('TF-IDF usa unigramas/bigramas, min_df=2 y hasta 100000 características; regresión logística L2 con optimización SGD. Todos los modelos neuronales usan EmbeddingBag(mean) → MLP de 128 unidades, ReLU, dropout 0.2 y salida de cuatro clases. Adam 0.001, batch 512, máximo 10 epochs y parada tras dos epochs sin mejorar F1 macro.'),
      p('La inicialización aleatoria tiene 100d. SGNS, gensim y GloVe se comparan congelados y ajustables, adaptando solo la entrada a su dimensión. Los OOV preentrenados parten de cero; congelados no aportan señal léxica y ajustables pueden aprender. Validación fija variante y checkpoint antes de evaluar test; se guardan hashes, predicciones y matrices.'),
      table(['Fracción','TF-IDF','Aleatorio','SGNS','Gensim','GloVe'],fractionrows),
      p('F1 macro en test por fracción. Cada columna usa la variante elegida exclusivamente con validación para esa fracción; no se cambió la selección después de ver test.'),
      table(['Modelo al 100%','Acc','P macro','R macro','F1 macro','Par entrenables','Min'],fullmetrics),
      p(f"El mayor F1 al 100% fue {winner['name']} ({percent(winner['metrics']['f1_macro'])}). La matriz con más errores fue {confusion['name']}; su mayor confusión dirigida fue {label[i]}→{label[j]} ({count} ejemplos). Las veinte matrices y curvas completas están guardadas en el notebook y archivos JSON/CSV."),
      p('Los embeddings congelados exigen menos parámetros entrenables; fine-tuning adapta dominio y OOV, pero requiere etiquetas. La referencia TF-IDF puede competir bien en noticias por su señal léxica y bigramas. La ventaja observada con poco entrenamiento se interpreta en la siguiente página, sin atribuir causalidad a una única semilla.')])
    time_sgns=sum(r['training_seconds'] for r in runs[best['config']['name']]); time_gen=sum(r['training_seconds'] for r in read('runs/gensim.json'))
    compare=[]
    for name,r in intrinsic.items():
        compare.append([name,r['dimension'],number(r['vocabulary_size']),rho(r['wordsim']['spearman']),percent(r['wordsim']['coverage']),rho(r['simlex']['spearman']),percent(r['simlex']['coverage']),percent(next(t for t in full if t['initialization']==name)['metrics']['f1_macro']),mins(time_sgns if name=='sgns' else time_gen) if name!='glove' else 'Preentrenado'])
    one={r['initialization']:r for r in test if r['fraction']==.01}
    gain=one['glove']['metrics']['f1_macro']-one['random']['metrics']['f1_macro']
    intrinsic_winner=max(intrinsic,key=lambda m:intrinsic[m]['simlex']['spearman'] or -2)
    pages.append([
      h('Conclusiones y referencias'),
      table(['Modelo','d','Vocab','ρ WS','Cob WS','ρ SL','Cob SL','F1 test','Min entren'],compare),
      p('WS=WordSim-353; SL=SimLex-999. Las correlaciones finales usan pares cubiertos por cada vocabulario, por lo que se muestra cobertura. SimLex se evaluó después de cerrar la selección de embeddings. Los tiempos de SGNS/gensim suman tres epochs de la configuración comparada, sin descarga ni evaluación.'),
      p('OOV de tokens en test AG News: '+', '.join(f"{name} {percent(read('news_split.json')['oov'][name]['test']['unknown_fraction'])}" for name in ['sgns','gensim','glove'])+'. SGNS y gensim usan '+number(best['corpus_tokens'])+' tokens normalizados; GloVe usa 6000000000. El hardware medido y el publicado se distinguen en la página 2.'),
      p(f"La configuración SGNS elegida fue {best['config']['name']} en epoch {best['epoch']}. {intrinsic_winner} obtuvo la mayor correlación SimLex observada. La tabla muestra que calidad intrínseca y F1 de clasificación deben examinarse por separado: son tareas con objetivos y distribuciones distintos."),
      p(f"Con 1% de etiquetas, GloVe obtuvo F1 {percent(one['glove']['metrics']['f1_macro'])} frente a {percent(one['random']['metrics']['f1_macro'])} del embedding aleatorio: diferencia {gain*100:+.2f} puntos porcentuales. Esta diferencia cuantifica la utilidad observada de la inicialización preentrenada con pocas etiquetas; no demuestra superioridad universal."),
      p('Subsampling redujo pares frecuentes redundantes. Dimensión, ventana y negativos cambian capacidad y costo; su efecto observado por configuración se analiza en el notebook. SGNS/gensim difieren en actualización; GloVe dispone de muchos más tokens. Esta comparación no aísla el efecto de arquitectura.'),
      p('Limitaciones: tres epochs, una semilla, sin intervalos de confianza; embeddings estáticos y promedio sin orden ni negación; tokenización y cobertura diferentes. Futuro: varias semillas, más entrenamiento y modelos sensibles al contexto.'),
      p('Verificación: pérdida y gradientes contra entropía cruzada binaria, fronteras de contexto, exclusión de consultas y coincidencia con gensim; corpus mínimo y hashes, subconjuntos anidados, selección bloqueada y una evaluación test por modelo. El notebook se ejecutó desde kernel limpio. Los checkpoints y registros reales se recuperaron de Colab.'),
      h('Referencias'),
      p('Mikolov et al. (2013). Distributed Representations of Words and Phrases and their Compositionality. NeurIPS. Pennington, Socher y Manning (2014). GloVe: Global Vectors for Word Representation. EMNLP. https://nlp.stanford.edu/pubs/glove.pdf'),
      p('PyTorch 2.8: Embedding y EmbeddingBag. Gensim 4.4: Word2Vec y KeyedVectors. Fuentes de datos: Salesforce/WikiText-103 raw y fancyzhx/AG News en Hugging Face; WordSim-353, SimLex-999 y analogías en los recursos de gensim.'),
      p('Entregables: notebook comentado, PDF, Word editable, código y dependencias, métricas, figuras, predicciones y mejores vectores SGNS .kv en la release v1.0-lab7.'),
      p('Repositorio de la entrega: '+URL)])
    return pages

def build_word(pages):
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.65)
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.03
    for name,size in [('Title',18),('Heading 1',14)]:
        style=doc.styles[name]; style.font.name='Arial'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(0); style.paragraph_format.space_after=Pt(7)
    header=sec.header.paragraphs[0]; header.text='Laboratorio 7  |  Ian Cumes 23236'; header.style='Caption'; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(0,0,0)
    footer=sec.footer.paragraphs[0]; footer.alignment=2
    footer.add_run('Página ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    for index,blocks in enumerate(pages):
        if index: doc.add_page_break()
        else:
            doc.add_paragraph('NLP y embeddings',style='Title')
            doc.add_paragraph('Laboratorio 7 · Ian Cumes · carné 23236 · entrega individual')
        for block in blocks:
            if block[0]=='h': doc.add_paragraph(block[1],style='Heading 1')
            elif block[0]=='p': doc.add_paragraph(block[1])
            elif block[0]=='image': doc.add_picture(str(ART/'figures'/block[1]),width=Inches(7.2))
            else:
                _,headers,rows=block
                t=doc.add_table(rows=1,cols=len(headers)); t.autofit=True
                for cell,text in zip(t.rows[0].cells,headers): cell.text=str(text)
                for row in rows:
                    for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
                for ri,row in enumerate(t.rows):
                    pr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); pr.append(cant)
                    if ri==0:
                        repeat=OxmlElement('w:tblHeader'); pr.append(repeat)
                    for cell in row.cells:
                        tcpr=cell._tc.get_or_add_tcPr()
                        borders=OxmlElement('w:tcBorders')
                        for edge in ['top','left','bottom','right']:
                            item=OxmlElement('w:'+edge); item.set(qn('w:val'),'single'); item.set(qn('w:sz'),'4'); item.set(qn('w:color'),'D7D7D7'); borders.append(item)
                        tcpr.append(borders)
                        if ri==0:
                            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F1F1F1'); tcpr.append(shade)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after=Pt(2); paragraph.paragraph_format.space_before=Pt(2)
                            for run in paragraph.runs: run.font.size=Pt(8.5); run.bold=ri==0
    doc.core_properties.author='Ian Cumes'; doc.core_properties.title='Laboratorio 7 NLP y embeddings'
    doc.core_properties.subject='SGNS propio, evaluación de embeddings y clasificación AG News'
    doc.save(OUT/'Laboratorio7_Ian_Cumes_23236.docx')

def build_pdf(pages):
    candidates=[(Path('C:/Windows/Fonts/arial.ttf'),Path('C:/Windows/Fonts/arialbd.ttf')),
                (Path('/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'),Path('/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf')),
                (Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))]
    regular,bold=next((a,b) for a,b in candidates if a.exists() and b.exists())
    pdfmetrics.registerFont(TTFont('ArialLab',str(regular)))
    pdfmetrics.registerFont(TTFont('ArialLabBold',str(bold)))
    style=ParagraphStyle('body',fontName='ArialLab',fontSize=10,leading=12.5,spaceAfter=6)
    heading=ParagraphStyle('heading',parent=style,fontName='ArialLabBold',fontSize=14,leading=17,spaceAfter=8)
    small=ParagraphStyle('cell',parent=style,fontSize=8,leading=10,spaceAfter=0)
    story=[]
    for index,blocks in enumerate(pages):
        if index: story.append(PageBreak())
        else:
            story.append(Paragraph('NLP y embeddings',ParagraphStyle('title',parent=heading,fontSize=18,leading=21)))
            story.append(Paragraph('Laboratorio 7 · Ian Cumes · carné 23236 · entrega individual',style))
        for block in blocks:
            if block[0]=='h': story.append(Paragraph(html.escape(block[1]),heading))
            elif block[0]=='p':
                text=html.escape(block[1]).replace(URL,f'<link href="{URL}">{URL}</link>')
                story.append(Paragraph(text,style))
            elif block[0]=='image':
                from reportlab.lib.utils import ImageReader
                imagepath=str(ART/'figures'/block[1]); width,height=ImageReader(imagepath).getSize()
                scale=min(7.2*72/width,block[2]*72/height)
                story.append(Image(imagepath,width=width*scale,height=height*scale))
            else:
                _,headers,rows=block
                data=[[Paragraph(html.escape(str(x)),small) for x in row] for row in [headers]+rows]
                # Adaptar anchuras a la longitud de las columnas, conservando mínimos legibles.
                lengths=[max(len(str(row[i])) for row in [headers]+rows) for i in range(len(headers))]
                weights=[max(7,min(30,n)) for n in lengths]
                widths=[7.2*72*w/sum(weights) for w in weights]
                t=Table(data,colWidths=widths,repeatRows=1)
                t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#F1F1F1')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#D7D7D7')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
                story.extend([t,Spacer(1,6)])
    def header(canvas,document):
        canvas.setFont('ArialLab',8); canvas.setFillColor(colors.HexColor('#555555'))
        canvas.drawString(46.8,766,'Laboratorio 7  |  Ian Cumes 23236')
        canvas.drawRightString(565,25,f'Página {document.page}')
    path=OUT/'Laboratorio7_Ian_Cumes_23236.pdf'
    SimpleDocTemplate(str(path),pagesize=(612,792),leftMargin=46.8,rightMargin=46.8,topMargin=42,bottomMargin=40,title='Laboratorio 7 NLP y embeddings',author='Ian Cumes').build(story,onFirstPage=header,onLaterPages=header)
    assert len(PdfReader(path).pages)<=5,'El PDF excede cinco páginas; ajustar composición.'

if __name__=='__main__':
    OUT.mkdir(exist_ok=True)
    pages=build_content()
    # Una sola fuente de contenido para garantizar equivalencia PDF/Word.
    (OUT/'contenido_informe.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
    build_word(pages); build_pdf(pages)
    print('PDF y Word generados a partir del mismo contenido.')
